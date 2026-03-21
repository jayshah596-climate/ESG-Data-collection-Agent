"""Module 8: Word Report Generation - Professional ESG summary report."""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from esg_agent.models import ESGDataset, ESGField


def _format_value(field: ESGField) -> str:
    """Format an ESGField value for display."""
    if field.is_missing():
        return "Not disclosed"
    if isinstance(field.value, list):
        return ", ".join(str(v) for v in field.value)
    val = str(field.value)
    if field.unit and field.unit not in val:
        val += f" {field.unit}"
    return val


def _add_table(doc: Document, headers: list, rows: list) -> None:
    """Add a formatted table to the document."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Headers
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True

    # Data rows
    for row_idx, row_data in enumerate(rows):
        for col_idx, cell_value in enumerate(row_data):
            table.rows[row_idx + 1].cells[col_idx].text = str(cell_value)

    doc.add_paragraph()  # spacing


def export_word(dataset: ESGDataset, output_path: str) -> str:
    """Generate a professional ESG summary report in Word format."""
    doc = Document()

    # Title
    title = doc.add_heading("ESG Data Summary Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    company_name = _format_value(dataset.company_profile.name)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(company_name)
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(68, 114, 196)

    reporting_year = _format_value(dataset.company_profile.reporting_year)
    year_para = doc.add_paragraph()
    year_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = year_para.add_run(f"Reporting Year: {reporting_year}")
    run.font.size = Pt(12)
    run.font.italic = True

    doc.add_paragraph()

    # Executive Summary
    doc.add_heading("1. Executive Summary", level=1)
    completeness = dataset.data_quality.completeness_pct
    missing_count = len(dataset.data_quality.fields_missing)

    exec_text = (
        f"This report presents the ESG data profile for {company_name}. "
        f"The automated data collection process achieved {completeness}% data completeness, "
        f"with {missing_count} data fields requiring manual input or further disclosure. "
    )

    if completeness >= 70:
        exec_text += "The overall disclosure level is strong, indicating mature ESG reporting practices."
    elif completeness >= 40:
        exec_text += "The disclosure level is moderate, with several areas requiring additional data."
    else:
        exec_text += "The disclosure level is limited. Significant manual data collection is recommended."

    doc.add_paragraph(exec_text)

    # Company Profile
    doc.add_heading("2. Company Profile", level=1)
    profile = dataset.company_profile
    _add_table(doc,
        ["Field", "Value"],
        [
            ["Company Name", _format_value(profile.name)],
            ["Industry", _format_value(profile.industry)],
            ["Headquarters", _format_value(profile.headquarters)],
            ["Reporting Year", _format_value(profile.reporting_year)],
        ],
    )

    # Emissions Overview
    doc.add_heading("3. Emissions Overview", level=1)
    ghg = dataset.ghg_emissions
    doc.add_paragraph(
        "The following table summarises the organisation's greenhouse gas emissions across all scopes."
    )
    _add_table(doc,
        ["Scope", "Emissions", "Confidence"],
        [
            ["Scope 1 (Direct)", _format_value(ghg.scope_1), ghg.scope_1.confidence or "N/A"],
            ["Scope 2 (Indirect - Energy)", _format_value(ghg.scope_2), ghg.scope_2.confidence or "N/A"],
            ["Scope 3 (Value Chain)", _format_value(ghg.scope_3), ghg.scope_3.confidence or "N/A"],
            ["Total GHG Emissions", _format_value(ghg.total), ghg.total.confidence or "N/A"],
        ],
    )

    # Energy
    doc.add_heading("4. Energy Consumption", level=1)
    energy = dataset.energy
    _add_table(doc,
        ["Metric", "Value", "Confidence"],
        [
            ["Total Energy Consumption", _format_value(energy.total_consumption), energy.total_consumption.confidence or "N/A"],
            ["Renewable Energy Share", _format_value(energy.renewable_energy_pct), energy.renewable_energy_pct.confidence or "N/A"],
        ],
    )

    # Climate Targets
    doc.add_heading("5. Climate Targets", level=1)
    targets = dataset.climate_targets
    _add_table(doc,
        ["Target", "Details", "Confidence"],
        [
            ["Net Zero Target Year", _format_value(targets.net_zero_year), targets.net_zero_year.confidence or "N/A"],
            ["Interim Targets", _format_value(targets.interim_targets), targets.interim_targets.confidence or "N/A"],
            ["Science-Based Targets", _format_value(targets.science_based_targets), targets.science_based_targets.confidence or "N/A"],
        ],
    )

    # Risk Assessment
    doc.add_heading("6. Climate Risk Assessment", level=1)
    risks = dataset.climate_risks

    doc.add_heading("Physical Risks", level=2)
    physical_val = _format_value(risks.physical_risks)
    doc.add_paragraph(f"Identified physical risks: {physical_val}")

    doc.add_heading("Transition Risks", level=2)
    transition_val = _format_value(risks.transition_risks)
    doc.add_paragraph(f"Identified transition risks: {transition_val}")

    # Governance
    doc.add_heading("7. Governance Insights", level=1)
    gov = dataset.governance
    _add_table(doc,
        ["Governance Area", "Status", "Confidence"],
        [
            ["Board Oversight", _format_value(gov.board_oversight), gov.board_oversight.confidence or "N/A"],
            ["ESG Committees", _format_value(gov.esg_committees), gov.esg_committees.confidence or "N/A"],
            ["Climate-Linked Incentives", _format_value(gov.climate_linked_incentives), gov.climate_linked_incentives.confidence or "N/A"],
        ],
    )

    # Scope 3 Breakdown
    doc.add_heading("8. Scope 3 Emissions Breakdown", level=1)
    s3 = dataset.scope3_breakdown
    _add_table(doc,
        ["Category", "Emissions", "Confidence"],
        [
            ["Purchased Goods & Services", _format_value(s3.purchased_goods_and_services), s3.purchased_goods_and_services.confidence or "N/A"],
            ["Transport & Distribution", _format_value(s3.transport_and_distribution), s3.transport_and_distribution.confidence or "N/A"],
            ["Use of Sold Products", _format_value(s3.use_of_sold_products), s3.use_of_sold_products.confidence or "N/A"],
            ["End-of-Life Treatment", _format_value(s3.end_of_life), s3.end_of_life.confidence or "N/A"],
        ],
    )

    # Data Gaps & Recommendations
    doc.add_heading("9. Data Gaps & Recommendations", level=1)
    missing = dataset.data_quality.fields_missing

    if missing:
        doc.add_paragraph(
            f"The following {len(missing)} data fields were not found in the available sources "
            "and require manual data collection or further company disclosure:"
        )
        for field_path in missing:
            doc.add_paragraph(f"  {field_path.replace('_', ' ').replace('.', ' > ').title()}", style="List Bullet")

        doc.add_paragraph()
        doc.add_paragraph(
            "Recommendation: Use the provided Excel template to manually input the missing data points. "
            "Engage with the company's sustainability team to obtain verified figures for incomplete fields."
        )
    else:
        doc.add_paragraph("All data fields have been populated. No significant data gaps identified.")

    # Data Quality Summary
    doc.add_heading("10. Data Quality Summary", level=1)
    doc.add_paragraph(f"Overall Data Completeness: {completeness}%")

    confidence_scores = dataset.data_quality.confidence_scores
    if confidence_scores:
        rows = []
        for section, scores in confidence_scores.items():
            section_label = section.replace("_", " ").title()
            rows.append([
                section_label,
                str(scores.get("high", 0)),
                str(scores.get("medium", 0)),
                str(scores.get("low", 0)),
                str(scores.get("missing", 0)),
            ])
        _add_table(doc, ["Section", "High", "Medium", "Low", "Missing"], rows)

    # Footer
    doc.add_paragraph()
    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Generated by ESG Data Collection & Intelligence System")
    run.font.size = Pt(9)
    run.font.italic = True
    run.font.color.rgb = RGBColor(128, 128, 128)

    doc.save(output_path)
    print(f"[INFO] Word report saved to {output_path}")
    return output_path
