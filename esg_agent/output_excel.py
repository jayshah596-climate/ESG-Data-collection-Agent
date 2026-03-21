"""Module 7: Excel Output Generation - Multi-sheet workbook with openpyxl."""

from dataclasses import fields as dc_fields
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from esg_agent.models import ESGDataset, ESGField


# Color coding for confidence levels
CONFIDENCE_FILLS = {
    "High": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),  # green
    "Medium": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),  # yellow
    "Low": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),  # red
}

HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def _format_value(value) -> str:
    """Format a value for display in Excel."""
    if value == "missing" or value is None:
        return "Missing"
    if isinstance(value, list):
        return ", ".join(str(v) for v in value)
    return str(value)


def _add_sheet(wb: Workbook, sheet_name: str, section_obj) -> None:
    """Add a sheet for an ESG section with Field, Value, Unit, Confidence, Source columns."""
    ws = wb.create_sheet(title=sheet_name)

    # Headers
    headers = ["Field", "Value", "Unit", "Confidence", "Source"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
        cell.border = THIN_BORDER

    # Data rows
    row = 2
    for f in dc_fields(section_obj):
        field_val = getattr(section_obj, f.name)
        if not isinstance(field_val, ESGField):
            continue

        # Field name (formatted)
        field_label = f.name.replace("_", " ").title()
        ws.cell(row=row, column=1, value=field_label).border = THIN_BORDER

        # Value
        cell_val = ws.cell(row=row, column=2, value=_format_value(field_val.value))
        cell_val.border = THIN_BORDER

        # Unit
        ws.cell(row=row, column=3, value=field_val.unit).border = THIN_BORDER

        # Confidence with color
        conf_cell = ws.cell(row=row, column=4, value=field_val.confidence or "N/A")
        conf_cell.border = THIN_BORDER
        if field_val.confidence in CONFIDENCE_FILLS:
            conf_cell.fill = CONFIDENCE_FILLS[field_val.confidence]

        # Source
        source_text = field_val.source[:200] if field_val.source else ""
        ws.cell(row=row, column=5, value=source_text).border = THIN_BORDER

        row += 1

    # Auto-adjust column widths
    for col in ws.columns:
        max_length = 0
        column_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[column_letter].width = min(max_length + 4, 60)


def export_excel(dataset: ESGDataset, output_path: str) -> str:
    """Export ESG dataset to a multi-sheet Excel workbook."""
    wb = Workbook()

    # Remove default sheet
    wb.remove(wb.active)

    # Add sheets for each section
    sections = [
        ("Company Profile", dataset.company_profile),
        ("GHG Emissions", dataset.ghg_emissions),
        ("Energy", dataset.energy),
        ("Climate Targets", dataset.climate_targets),
        ("Climate Risks", dataset.climate_risks),
        ("Governance", dataset.governance),
        ("Scope 3 Breakdown", dataset.scope3_breakdown),
    ]

    for sheet_name, section_obj in sections:
        _add_sheet(wb, sheet_name, section_obj)

    # Add Data Quality summary sheet
    ws_quality = wb.create_sheet(title="Data Quality")
    ws_quality.cell(row=1, column=1, value="Metric").font = HEADER_FONT
    ws_quality.cell(row=1, column=1).fill = HEADER_FILL
    ws_quality.cell(row=1, column=2, value="Value").font = HEADER_FONT
    ws_quality.cell(row=1, column=2).fill = HEADER_FILL

    ws_quality.cell(row=2, column=1, value="Data Completeness (%)")
    ws_quality.cell(row=2, column=2, value=dataset.data_quality.completeness_pct)

    ws_quality.cell(row=3, column=1, value="Missing Fields Count")
    ws_quality.cell(row=3, column=2, value=len(dataset.data_quality.fields_missing))

    ws_quality.cell(row=4, column=1, value="Missing Fields")
    ws_quality.cell(row=4, column=2, value=", ".join(dataset.data_quality.fields_missing))

    for col in ws_quality.columns:
        column_letter = col[0].column_letter
        ws_quality.column_dimensions[column_letter].width = 40

    wb.save(output_path)
    print(f"[INFO] Excel output saved to {output_path}")
    return output_path
